#!/usr/bin/env python3
"""Build the AjouNice perception-only midterm report from the official template.

The generated deck intentionally contains only the perception contribution so
the team lead can copy or merge the slides into the complete team report.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


NAVY = RGBColor(23, 48, 84)
BLUE = RGBColor(31, 119, 180)
CYAN = RGBColor(36, 166, 183)
GREEN = RGBColor(36, 145, 96)
ORANGE = RGBColor(235, 134, 52)
PURPLE = RGBColor(123, 84, 170)
RED = RGBColor(196, 62, 62)
LIGHT = RGBColor(241, 246, 250)
MID = RGBColor(219, 230, 239)
DARK = RGBColor(35, 42, 50)
GRAY = RGBColor(96, 106, 116)
WHITE = RGBColor(255, 255, 255)


def remove_shape(shape) -> None:
    element = shape._element
    element.getparent().remove(element)


def replace_text(slide, old: str, new: str) -> None:
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new)


def set_textbox(
    slide,
    x,
    y,
    w,
    h,
    text: str,
    *,
    size: float = 18,
    bold: bool = False,
    color: RGBColor = DARK,
    align=PP_ALIGN.LEFT,
    font: str = "Noto Sans CJK KR",
    margin: float = 0.05,
):
    shape = slide.shapes.add_textbox(x, y, w, h)
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return shape


def add_box(
    slide,
    x,
    y,
    w,
    h,
    title: str,
    body: str,
    *,
    accent: RGBColor = BLUE,
    fill: RGBColor = LIGHT,
    title_size: float = 17,
    body_size: float = 13,
):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    box.fill.solid()
    box.fill.fore_color.rgb = fill
    box.line.color.rgb = accent
    box.line.width = Pt(1.4)

    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, Inches(0.09), h)
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = accent
    stripe.line.fill.background()

    set_textbox(
        slide,
        x + Inches(0.18),
        y + Inches(0.09),
        w - Inches(0.28),
        Inches(0.30),
        title,
        size=title_size,
        bold=True,
        color=accent,
    )
    body_shape = slide.shapes.add_textbox(
        x + Inches(0.18),
        y + Inches(0.43),
        w - Inches(0.30),
        h - Inches(0.50),
    )
    tf = body_shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for idx, line in enumerate(body.split("\n")):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.level = 0
        p.space_after = Pt(4)
        p.font.name = "Noto Sans CJK KR"
        p.font.size = Pt(body_size)
        p.font.color.rgb = DARK
    return box


def add_metric(slide, x, y, w, label: str, value: str, color: RGBColor):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, Inches(0.77))
    box.fill.solid()
    box.fill.fore_color.rgb = WHITE
    box.line.color.rgb = color
    box.line.width = Pt(1.3)
    set_textbox(
        slide, x + Inches(0.08), y + Inches(0.07), w - Inches(0.16), Inches(0.23),
        label, size=10.5, bold=True, color=GRAY, align=PP_ALIGN.CENTER,
    )
    set_textbox(
        slide, x + Inches(0.08), y + Inches(0.29), w - Inches(0.16), Inches(0.36),
        value, size=18, bold=True, color=color, align=PP_ALIGN.CENTER,
    )


def add_pipeline_stage(slide, x, y, w, title, detail, color):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, Inches(1.12))
    box.fill.solid()
    box.fill.fore_color.rgb = color
    box.line.color.rgb = color
    set_textbox(
        slide, x + Inches(0.05), y + Inches(0.12), w - Inches(0.10), Inches(0.31),
        title, size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
    )
    set_textbox(
        slide, x + Inches(0.08), y + Inches(0.48), w - Inches(0.16), Inches(0.46),
        detail, size=10.5, color=WHITE, align=PP_ALIGN.CENTER,
    )
    return box


def add_arrow(slide, x1, y1, x2, y2):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    line.line.color.rgb = NAVY
    line.line.width = Pt(2)
    line.line.end_arrowhead = True
    return line


def build(template: Path, output: Path) -> None:
    prs = Presentation(str(template))
    if len(prs.slides) != 6:
        raise ValueError(f"expected six-slide official template, got {len(prs.slides)}")

    # Cover and table of contents retain the official visual theme.
    replace_text(prs.slides[0], "학교명", "아주대학교")
    replace_text(prs.slides[0], "팀명", "AjouNice")
    body = prs.slides[1].shapes[1].text_frame
    body.clear()
    for idx, text in enumerate(
        [
            "개발 현황 — 센서 기반 3D 객체인지·추적·경로예측",
            "알고리즘 소개 — PointPainting → EKF → Lanelet Map-CV",
            "문제점 및 개선 필요 사항 — online domain gap·heading 기준",
            "추후 개발 일정 — 판단 연동·실주행 검증·강건화",
        ]
    ):
        p = body.paragraphs[0] if idx == 0 else body.add_paragraph()
        p.text = text
        p.font.name = "Noto Sans CJK KR"
        p.font.size = Pt(20)
        p.space_after = Pt(16)

    # Slide 3: current status.
    s = prs.slides[2]
    add_box(
        s, Inches(0.55), Inches(1.70), Inches(3.92), Inches(2.18),
        "3D Object Detection",
        "• 4-camera semantic + VLP-16 LiDAR 융합\n"
        "• Vehicle / Pedestrian / Obstacle 3-class\n"
        "• PointPainting + PointPillars single-frame\n"
        "• MORAI UDP 센서 입력 기반 실시간 추론",
        accent=BLUE,
    )
    add_box(
        s, Inches(4.70), Inches(1.70), Inches(3.92), Inches(2.18),
        "Tracking",
        "• 표준 DetectedObjects → TrackedObjects 계약\n"
        "• EKF 기반 위치·속도 상태 추정\n"
        "• class별 association 및 track 관리\n"
        "• detector velocity 없이 tracker가 속도 추정",
        accent=GREEN,
    )
    add_box(
        s, Inches(8.85), Inches(1.70), Inches(3.92), Inches(2.18),
        "Prediction",
        "• Lanelet2 v60 기반 Map-CV\n"
        "• 분기별 multi-path hypothesis 생성\n"
        "• 1·2·3초 predicted path 제공\n"
        "• planning 표준 PredictedObjects 출력",
        accent=PURPLE,
    )
    add_box(
        s, Inches(0.55), Inches(4.15), Inches(5.95), Inches(1.50),
        "데이터·모델 현황",
        "• MORAI class-centric Base3-Static 10,902 frame 구축\n"
        "• S1g E37을 현재 안정적인 실시간 시연 기준으로 사용\n"
        "• Base3 계열은 detector 표현력 개선 후 calibration 검증 중",
        accent=ORANGE,
        body_size=12.5,
    )
    add_box(
        s, Inches(6.75), Inches(4.15), Inches(6.02), Inches(1.50),
        "중간점검 시연 범위",
        "• 고주로에서 OD → Tracking → Prediction을 실시간 시각화\n"
        "• MORAI 화면과 RViz Lanelet overlay를 동시에 녹화\n"
        "• 판단·제어는 개발 중이므로 인지 pipeline 단독 시연",
        accent=CYAN,
        body_size=12.5,
    )

    # Slide 4: architecture and measured online results.
    s = prs.slides[3]
    for shape in list(s.shapes)[2:]:
        remove_shape(shape)
    set_textbox(
        s, Inches(0.55), Inches(1.26), Inches(4.0), Inches(0.36),
        "알고리즘 (AI 모델) 소개", size=17, bold=True, color=DARK,
    )
    stages = [
        ("센서 입력", "4 Camera\nVLP-16 LiDAR", BLUE),
        ("Semantic Painting", "LRASPP 4-class\npoint feature 결합", CYAN),
        ("3D OD", "PointPillars\n3D box + class", ORANGE),
        ("Tracking", "EKF\nposition + velocity", GREEN),
        ("Map-CV", "Lanelet2\n1·2·3 s paths", PURPLE),
    ]
    x0, y, w, gap = Inches(0.52), Inches(1.82), Inches(2.16), Inches(0.38)
    boxes = []
    for idx, (title, detail, color) in enumerate(stages):
        x = x0 + idx * (w + gap)
        boxes.append(add_pipeline_stage(s, x, y, w, title, detail, color))
        if idx:
            prev = boxes[idx - 1]
            add_arrow(
                s,
                prev.left + prev.width + Inches(0.05),
                y + Inches(0.56),
                x - Inches(0.05),
                y + Inches(0.56),
            )
    set_textbox(
        s, Inches(0.55), Inches(3.18), Inches(12.2), Inches(0.33),
        "동일한 standardized object message로 검출기 교체가 가능하며, 현재 수치는 S1g E37·score 0.40의 실주행 rosbag 평가임",
        size=11.5, color=GRAY, align=PP_ALIGN.CENTER,
    )
    add_metric(s, Inches(0.58), Inches(3.73), Inches(2.75), "로터리 OD center mAP", "0.720", BLUE)
    add_metric(s, Inches(3.45), Inches(3.73), Inches(2.75), "고주로 OD center mAP", "0.803", BLUE)
    add_metric(s, Inches(6.32), Inches(3.73), Inches(2.75), "로터리 Tracking F1", "0.960", GREEN)
    add_metric(s, Inches(9.19), Inches(3.73), Inches(2.75), "고주로 Tracking F1", "0.953", GREEN)
    add_box(
        s, Inches(0.58), Inches(4.72), Inches(5.62), Inches(1.23),
        "로터리 Map-CV prediction p50",
        "1초 0.464 m  |  2초 0.727 m  |  3초 1.096 m",
        accent=PURPLE, body_size=13.5,
    )
    add_box(
        s, Inches(6.32), Inches(4.72), Inches(5.62), Inches(1.23),
        "고주로 Map-CV prediction p50",
        "1초 0.491 m  |  2초 0.683 m  |  3초 0.956 m",
        accent=PURPLE, body_size=13.5,
    )

    # Slide 5: limitations and improvements.
    s = prs.slides[4]
    add_box(
        s, Inches(0.55), Inches(1.73), Inches(3.92), Inches(3.55),
        "Online domain gap",
        "• 학습용 SaveSensorData와 실제 UDP 입력의 카메라·LiDAR 시각 차이\n"
        "• 고주로·로터리의 거리·가림·관측각 변화\n"
        "• offline validation 대비 실주행 confidence calibration 차이\n\n"
        "개선: 실제 rosbag slice 평가, class별 threshold, hard-negative mining",
        accent=BLUE,
        body_size=13,
    )
    add_box(
        s, Inches(4.70), Inches(1.73), Inches(3.92), Inches(3.55),
        "Localization heading 기준",
        "• Scenario/map yaw와 ROS heading 사이 약 1.347° 편차\n"
        "• K-City UTM Zone 52 자오선 수렴각과 약 0.001° 이내 일치\n"
        "• 변환 누락 시 30 m 전방에서 약 0.71 m 횡오차\n\n"
        "개선: true north → grid north 변환 계약 명확화 및 단일 적용",
        accent=RED,
        body_size=13,
    )
    add_box(
        s, Inches(8.85), Inches(1.73), Inches(3.92), Inches(3.55),
        "동적 객체·속도",
        "• single-frame OD의 직접 velocity 회귀는 관측 한계가 큼\n"
        "• 현재 production은 EKF history로 속도를 추정\n"
        "• 3-sweep은 ghosting·online 입력 정합 문제가 남음\n\n"
        "개선: detector velocity 유무 A/B, tracking AVE 및 1·2·3초 FDE 검증",
        accent=ORANGE,
        body_size=13,
    )
    set_textbox(
        s, Inches(0.65), Inches(5.52), Inches(12.0), Inches(0.48),
        "대회 제출 시연은 허용된 UDP 센서만 사용하며, /Object_topic GT 및 MORAI BBox overlay는 사용하지 않음",
        size=13, bold=True, color=RED, align=PP_ALIGN.CENTER,
    )

    # Slide 6: perception roadmap.
    s = prs.slides[5]
    add_box(
        s, Inches(0.65), Inches(1.72), Inches(3.78), Inches(3.75),
        "1단계  |  대회 입력 안정화",
        "• 대회용 UDP sensor rate·timestamp 검증\n"
        "• heading reference 공식 확인 및 localization 변환 고정\n"
        "• 실시간 latency·drop·TF 상태 점검\n"
        "• 고주로/로터리 반복 시연 안정화",
        accent=BLUE,
        body_size=13.5,
    )
    add_box(
        s, Inches(4.78), Inches(1.72), Inches(3.78), Inches(3.75),
        "2단계  |  인지 강건화",
        "• 전체 경로 rosbag 기반 false positive 수집\n"
        "• 중앙분리대·가드레일 hard-negative mining\n"
        "• Pedestrian·Obstacle confidence calibration\n"
        "• mission-relevant slice별 회귀 검증",
        accent=GREEN,
        body_size=13.5,
    )
    add_box(
        s, Inches(8.91), Inches(1.72), Inches(3.78), Inches(3.75),
        "3단계  |  판단·제어 연동",
        "• PredictedObjects와 행동판단 interface 통합\n"
        "• 합류·추월·횡단보도·GPS 음영 미션 검증\n"
        "• prediction-aware planning A/B\n"
        "• 전체 코스 end-to-end 반복 주행",
        accent=PURPLE,
        body_size=13.5,
    )
    set_textbox(
        s, Inches(0.75), Inches(5.76), Inches(11.85), Inches(0.40),
        "중간점검 이후 우선순위: localization 기준 확정 → 실주행 인지 안정화 → 판단·제어 통합",
        size=14.5, bold=True, color=NAVY, align=PP_ALIGN.CENTER,
    )

    # Metadata and save.
    prs.core_properties.title = "AjouNice 중간점검 인지 개발 현황"
    prs.core_properties.subject = "PointPainting 3D OD, EKF tracking, Lanelet Map-CV prediction"
    prs.core_properties.author = "AjouNice"
    prs.core_properties.comments = "Perception-only draft generated from the official MORAI template."
    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    if not args.template.is_file():
        raise FileNotFoundError(args.template)
    build(args.template, args.output)
    print(args.output)


if __name__ == "__main__":
    main()
